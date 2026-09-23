from io import BytesIO

import docx
import openpyxl
import pymupdf
import pytest
from conftest import edition_path

from la_rp_peace.enums import DocFormat
from la_rp_peace.ingestion.extract import ExtractionError, UnsupportedFormatError, detect_format, extract
from la_rp_peace.ingestion.extract.types import Extraction


def _extract_file(edition: int, suffix: str) -> Extraction:
    path = edition_path(edition, suffix)
    data = path.read_bytes()
    return extract(detect_format(path.name, data), data)


def _assert_offsets_consistent(extraction: Extraction) -> None:
    for block in extraction.blocks:
        assert extraction.original_text[block.start : block.end] == block.text
    for span in extraction.source_map:
        assert 0 <= span.start < span.end <= len(extraction.original_text)


@pytest.mark.parametrize("suffix", [".docx", ".pdf"])
def test_blocks_and_source_map_point_into_original_text(suffix: str) -> None:
    extraction = _extract_file(9, suffix)

    _assert_offsets_consistent(extraction)
    assert "3.4. БВА состоит из следующих структурных подразделений:" in extraction.original_text


def test_docx_locations_are_body_positions() -> None:
    extraction = _extract_file(9, ".docx")

    block = next(block for block in extraction.blocks if block.text.startswith("а. Департамент ИТ-аудита"))

    assert block.location == {"paragraph": 103}
    assert block.style == "Normal"


def test_pdf_lines_keep_their_pages_and_page_numbers_are_kept() -> None:
    extraction = _extract_file(9, ".pdf")

    item = next(block for block in extraction.blocks if block.text.startswith("а. Департамент ИТ-аудита"))
    page_numbers = [block for block in extraction.blocks if block.style == "page_number"]

    assert item.location == {"page": 6}
    assert [block.text for block in page_numbers] == [str(page) for page in range(1, 26)]


@pytest.mark.parametrize("edition", [8, 9])
def test_same_words_from_docx_and_pdf(edition: int) -> None:
    from_docx = _extract_file(edition, ".docx")
    from_pdf = _extract_file(edition, ".pdf")

    pdf_words = " ".join(block.text for block in from_pdf.blocks if block.style != "page_number").split()

    assert pdf_words == from_docx.original_text.split()


def _synthetic_pdf() -> bytes:
    pdf = pymupdf.open()
    first = pdf.new_page()
    first.insert_text((72, 72), "1. General provisions", fontname="hebo", fontsize=12)
    first.insert_text((72, 100), "1.1. The unit prepares the cost-", fontsize=11)
    first.insert_text((72, 114), "benefit plan", fontsize=11)
    first.insert_text((72, 128), "- a dash that only wrapped", fontsize=11)
    first.insert_text((300, 770), "1", fontsize=9)
    second = pdf.new_page()
    second.insert_text((72, 72), "and continues on the next page.", fontsize=11)
    second.insert_text((72, 100), "1.2. Second clause.", fontsize=11)
    second.insert_text((72, 114), "- a real dash item;", fontsize=11)
    second.insert_text((72, 150), "2. Rights and", fontname="hebo", fontsize=12)
    second.insert_text((72, 164), "duties", fontname="hebo", fontsize=12)
    second.insert_text((72, 200), "Separate bold heading", fontname="hebo", fontsize=12)
    data: bytes = pdf.tobytes()
    return data


def test_pdf_lines_are_rebuilt_into_paragraphs() -> None:
    extraction = extract(DocFormat.PDF, _synthetic_pdf())

    assert [(block.text, block.style, block.location) for block in extraction.blocks] == [
        ("1. General provisions", "bold", {"page": 1}),
        (
            "1.1. The unit prepares the cost-benefit plan - a dash that only wrapped and continues on the next page.",
            None,
            {"page": 1},
        ),
        ("1", "page_number", {"page": 1}),
        ("1.2. Second clause.", None, {"page": 2}),
        ("- a real dash item;", None, {"page": 2}),
        ("2. Rights and duties", "bold", {"page": 2}),
        ("Separate bold heading", "bold", {"page": 2}),
    ]
    _assert_offsets_consistent(extraction)
    wrapped = extraction.blocks[1]
    pages = [span.location["page"] for span in extraction.source_map if wrapped.start <= span.start < wrapped.end]
    assert pages == [1, 1, 1, 2]


def test_docx_table_rows_have_cells() -> None:
    document = docx.Document()
    document.add_paragraph("Штатное расписание")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Подразделение", "Руководитель"
    table.cell(1, 0).text, table.cell(1, 1).text = "ДИТААД", "Директор ДИТААД"
    buffer = BytesIO()
    document.save(buffer)

    extraction = extract(DocFormat.DOCX, buffer.getvalue())

    row = extraction.blocks[2]
    assert row.style == "table"
    assert row.text == "ДИТААД\tДиректор ДИТААД"
    assert [extraction.original_text[cell.start : cell.end] for cell in row.cells] == ["ДИТААД", "Директор ДИТААД"]
    assert row.cells[1].location == {"paragraph": 1, "table_row": 2, "column": 2}


def test_xlsx_cells_carry_sheet_row_and_column() -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Оргструктура"
    sheet.append(["Подразделение", "Руководитель"])
    sheet.append([None, None])
    sheet.append(["ДИТААД", "Директор ДИТААД"])
    buffer = BytesIO()
    workbook.save(buffer)

    extraction = extract(DocFormat.XLSX, buffer.getvalue())

    assert [block.location for block in extraction.blocks] == [
        {"sheet": "Оргструктура", "row": 1},
        {"sheet": "Оргструктура", "row": 3},
    ]
    assert extraction.blocks[1].cells[1].location == {"sheet": "Оргструктура", "row": 3, "column": "B"}
    _assert_offsets_consistent(extraction)


def test_file_properties_are_kept_with_their_source() -> None:
    extraction = _extract_file(9, ".pdf")

    assert all(entry["source"] == "pdf.metadata" for entry in extraction.file_metadata.values())


@pytest.mark.parametrize(
    ("filename", "data", "message"),
    [
        ("old.doc", b"\xd0\xcf\x11\xe0", "пересохраните файл как .docx"),
        ("notes.txt", b"hello", "Поддерживаются только"),
        ("fake.pdf", b"PK\x03\x04", "не соответствует формату"),
    ],
)
def test_unsupported_files_are_rejected(filename: str, data: bytes, message: str) -> None:
    with pytest.raises(UnsupportedFormatError, match=message):
        detect_format(filename, data)


def test_corrupt_file_raises_extraction_error() -> None:
    with pytest.raises(ExtractionError, match="Не удалось прочитать"):
        extract(DocFormat.DOCX, b"PK\x03\x04garbage")


def test_pdf_without_text_raises_extraction_error() -> None:
    pdf = pymupdf.open()
    pdf.new_page()

    with pytest.raises(ExtractionError, match="не найден текст"):
        extract(DocFormat.PDF, pdf.tobytes())
