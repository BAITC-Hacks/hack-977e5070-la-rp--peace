"""Extract paragraphs, table rows and file properties from Word (.docx) documents."""

from datetime import datetime
from io import BytesIO
from typing import Any

import docx
from docx.table import Table
from docx.text.paragraph import Paragraph

from la_rp_peace.enums import DocFormat
from la_rp_peace.ingestion.extract.types import Extraction, Location, TextBuilder
from la_rp_peace.ingestion.numbering import normalize_text

_PROPERTY_SOURCE = "docx.core_properties"
_PROPERTIES = ("title", "author", "last_modified_by", "created", "modified", "revision")


def _file_metadata(document: Any) -> dict[str, Any]:
    properties = document.core_properties
    metadata: dict[str, Any] = {}
    for name in _PROPERTIES:
        value = getattr(properties, name)
        if value in (None, "", 0):
            continue
        serialised = value.isoformat() if isinstance(value, datetime) else value
        metadata[name] = {"value": serialised, "source": _PROPERTY_SOURCE}
    return metadata


def _add_table(builder: TextBuilder, table: Table, index: int) -> None:
    for row_number, row in enumerate(table.rows, start=1):
        cells: list[tuple[str, Location]] = []
        previous = None
        for column, cell in enumerate(row.cells, start=1):
            # Merged cells repeat the same cell object in every grid column they span.
            if cell._tc is previous:
                continue
            previous = cell._tc
            text = normalize_text(cell.text)
            if text:
                cells.append((text, {"paragraph": index, "table_row": row_number, "column": column}))
        if cells:
            builder.add_row(cells, {"paragraph": index, "table_row": row_number})


def extract_docx(data: bytes) -> Extraction:
    """Read paragraphs and table rows of a .docx file in document order.

    Args:
        data: Raw file bytes.

    Returns:
        The extraction; locations carry the 0-based body item index as ``paragraph``.
    """
    document = docx.Document(BytesIO(data))
    builder = TextBuilder()
    for index, item in enumerate(document.iter_inner_content()):
        if isinstance(item, Table):
            _add_table(builder, item, index)
            continue
        paragraph: Paragraph = item
        text = normalize_text(paragraph.text)
        if text:
            style = paragraph.style.name if paragraph.style is not None else None
            builder.add(text, {"paragraph": index}, style=style)
    return Extraction(
        source_format=DocFormat.DOCX,
        original_text=builder.text(),
        blocks=builder.blocks,
        source_map=builder.source_map,
        file_metadata=_file_metadata(document),
    )
