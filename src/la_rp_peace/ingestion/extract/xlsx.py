"""Extract rows and cells from Excel (.xlsx) workbooks."""

from datetime import datetime
from io import BytesIO
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter

from la_rp_peace.enums import DocFormat
from la_rp_peace.ingestion.extract.types import Extraction, Location, TextBuilder
from la_rp_peace.ingestion.numbering import normalize_text

_PROPERTY_SOURCE = "xlsx.properties"
_PROPERTIES = ("title", "creator", "lastModifiedBy", "created", "modified")


def _file_metadata(workbook: Any) -> dict[str, Any]:
    metadata: dict[str, Any] = {}
    for name in _PROPERTIES:
        value = getattr(workbook.properties, name, None)
        if value in (None, ""):
            continue
        serialised = value.isoformat() if isinstance(value, datetime) else value
        metadata[name] = {"value": serialised, "source": _PROPERTY_SOURCE}
    return metadata


def extract_xlsx(data: bytes) -> Extraction:
    """Turn every non-empty row of every sheet into a row block with cell ranges.

    Formulas are read as their cached values; macros are never executed.

    Args:
        data: Raw file bytes.

    Returns:
        The extraction; cell locations carry sheet, row and column letter.
    """
    workbook = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
    builder = TextBuilder()
    try:
        metadata = _file_metadata(workbook)
        for sheet in workbook.worksheets:
            for row_number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
                cells: list[tuple[str, Location]] = []
                for column, value in enumerate(values, start=1):
                    text = normalize_text(str(value)) if value is not None else ""
                    if text:
                        location = {"sheet": sheet.title, "row": row_number, "column": get_column_letter(column)}
                        cells.append((text, location))
                if cells:
                    builder.add_row(cells, {"sheet": sheet.title, "row": row_number})
    finally:
        workbook.close()
    return Extraction(
        source_format=DocFormat.XLSX,
        original_text=builder.text(),
        blocks=builder.blocks,
        source_map=builder.source_map,
        file_metadata=metadata,
    )
